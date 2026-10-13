"""
XLSX extraction.

For each sheet: look in the first rows for a unit statement ("All amounts in USD '000") and a
header row containing period labels. Sheets without a period header (e.g. a 36,000-row general
ledger dump) are skipped after reading only those first rows, which keeps large workbooks fast.
Each data row's first text cell is the label; numeric cells under period columns become items
and blank cells in an otherwise filled row become "missing" observations.
"""

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.domain.periods import parse_period
from app.domain.types import ExtractedItem
from app.domain.units import MISSING_TOKENS, Unit, decimals_from_excel_format, parse_unit
from app.parsers import LabelMapper, ParseResult

_HEADER_SCAN_ROWS = 15


def parse_xlsx(path, mapper: LabelMapper) -> ParseResult:
    result = ParseResult(items=[])
    filename = str(path).rsplit("/", 1)[-1]
    # read_only streams rows; data_only returns cached formula results instead of formulas.
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            _parse_sheet(ws, filename, mapper, result)
    finally:
        wb.close()
    return result


def _parse_sheet(ws, filename: str, mapper: LabelMapper, result: ParseResult) -> None:
    unit: Unit | None = None
    columns: dict[int, tuple[str, str]] = {}  # column index -> (period, header text)
    title = ""
    rows = ws.iter_rows()
    for r_idx, row in enumerate(rows, start=1):
        if not columns:
            if r_idx > _HEADER_SCAN_ROWS:
                return  # no period header near the top: not a comparable statement
            texts = [(c_idx, str(c.value)) for c_idx, c in enumerate(row, start=1) if c.value is not None]
            for _, t in texts:
                u = parse_unit(t)
                if u and unit is None:
                    unit = u
                if not title and len(t) > 10:
                    title = t
            periods = {c_idx: (p, t) for c_idx, t in texts if (p := parse_period(t))}
            if periods:
                columns = periods
            continue
        _parse_row(row, r_idx, columns, unit, title, ws.title, filename, mapper, result)


def _parse_row(row, r_idx, columns, unit, title, sheet, filename, mapper, result) -> None:
    cells = list(row)
    label = next(
        (str(c.value) for c in cells[: min(3, len(cells))] if isinstance(c.value, str) and c.value.strip()), None
    )
    if label is None:
        return
    segment_ctx = "segment" in sheet.lower() or "segment" in title.lower()
    metric = mapper(label, segment_ctx)
    period_cells = [(c_idx, cells[c_idx - 1] if c_idx - 1 < len(cells) else None) for c_idx in columns]
    has_values = any(c is not None and isinstance(c.value, int | float) for _, c in period_cells)
    if metric is None:
        if has_values:
            result.unmapped_labels.append(label)
        return
    if not has_values:
        return
    multiplier = 1.0 if (metric == "headcount" or unit is None) else unit.multiplier
    unit_label = "persons" if metric == "headcount" else (unit.label if unit else "units")
    for c_idx, cell in period_cells:
        period, header = columns[c_idx]
        ref = f"{get_column_letter(c_idx)}{r_idx}"
        location = {"doc": filename, "sheet": sheet, "cell": ref, "row": label, "column": header}
        value = getattr(cell, "value", None)
        if value is None or (isinstance(value, str) and value.strip().lower() in MISSING_TOKENS):
            result.items.append(
                ExtractedItem(metric, period, None, multiplier, label, "", unit_label, location, is_missing=True)
            )
            continue
        if not isinstance(value, int | float):
            continue
        decimals = decimals_from_excel_format(getattr(cell, "number_format", None))
        result.items.append(
            ExtractedItem(
                metric=metric,
                period=period,
                value=float(value) * multiplier,
                precision=multiplier * 10 ** (-decimals),
                raw_label=label,
                raw_value=str(value),
                unit_label=unit_label,
                location=location,
            )
        )
