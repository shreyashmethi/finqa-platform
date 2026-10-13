"""
Evidence rendering for the review UI.

* PDF: render one page to PNG and draw a box around the conflicting value(s) using the bounding
  boxes captured at extraction time.
* XLSX: return a small grid of cells around the referenced cell.
* CSV: return the referenced line with a few lines of context.
"""

import csv
import io
from functools import lru_cache

import pypdfium2 as pdfium
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import coordinate_from_string
from PIL import Image, ImageDraw

from app.parsers.pdfium_lock import PDFIUM_LOCK

_RENDER_SCALE = 1.6  # 72 dpi * 1.6 ~ 115 dpi: readable without huge images


@lru_cache(maxsize=64)
def _render_page(path: str, page: int, mtime: float) -> bytes:
    """Render a page once and cache the PNG (keyed by file mtime so edits invalidate the cache)."""
    with PDFIUM_LOCK:  # PDFium is not thread-safe
        return _render_locked(path, page)


def _render_locked(path: str, page: int) -> bytes:
    doc = pdfium.PdfDocument(path)
    try:
        if not 1 <= page <= len(doc):
            raise ValueError(f"page {page} out of range")
        pdf_page = doc[page - 1]
        bitmap = pdf_page.render(scale=_RENDER_SCALE)
        img = bitmap.to_pil()
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        bitmap.close()
        pdf_page.close()
        return buf.getvalue()
    finally:
        doc.close()


def page_png(path: str, page: int, boxes: list[tuple[list[float], str]], mtime: float) -> bytes:
    """
    PNG of a page with highlight boxes. ``boxes`` = [(bbox [x0, top, x1, bottom] in PDF points, colour)].
    """
    img = Image.open(io.BytesIO(_render_page(path, page, mtime))).convert("RGB")
    overlay = Image.new("RGBA", img.size, (0, 0, 0, 0))
    draw = ImageDraw.Draw(overlay)
    colours = {"red": (220, 38, 38), "green": (22, 163, 74), "blue": (37, 99, 235)}
    for bbox, colour in boxes:
        x0, top, x1, bottom = (v * _RENDER_SCALE for v in bbox)
        pad = 3
        rgb = colours.get(colour, colours["blue"])
        draw.rectangle([x0 - pad, top - pad, x1 + pad, bottom + pad], fill=(*rgb, 40), outline=(*rgb, 255), width=3)
    img = Image.alpha_composite(img.convert("RGBA"), overlay).convert("RGB")
    buf = io.BytesIO()
    img.save(buf, format="PNG", optimize=True)
    return buf.getvalue()


def xlsx_window(path: str, sheet: str, cell: str, radius: int = 4) -> dict:
    """Cells around ``cell`` (plus the header rows) for display; returns rows of strings."""
    col_letter, row = coordinate_from_string(cell)
    col = column_index_from_string(col_letter)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet]
        min_row, max_row = max(1, row - radius), row + radius
        max_col = max(col + 2, 8)
        grid = []
        # Always include the top rows (title, unit, header) for context.
        wanted = sorted(set(range(1, 5)) | set(range(min_row, max_row + 1)))
        for r_idx, values in enumerate(
            ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col, values_only=True), start=1
        ):
            if r_idx in wanted:
                grid.append({"row": r_idx, "cells": ["" if v is None else str(v) for v in values]})
        return {
            "sheet": sheet,
            "columns": [get_column_letter(i) for i in range(1, max_col + 1)],
            "rows": grid,
            "highlight": {"row": row, "col": col},
        }
    finally:
        wb.close()


def csv_window(path: str, line: int, radius: int = 3) -> dict:
    with open(path, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))
    lo, hi = max(2, line - radius), min(len(rows), line + radius)
    return {
        "header": rows[0] if rows else [],
        "rows": [{"line": i, "cells": rows[i - 1]} for i in range(lo, hi + 1)],
        "highlight": line,
    }
