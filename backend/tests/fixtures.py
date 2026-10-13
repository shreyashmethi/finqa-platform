"""
Builds a tiny, self-contained document set for tests: a 2-page annual report PDF, a board deck PDF,
management accounts XLSX and a KPI CSV for a fictional company "T001" (USD, FY2025).

Known content:
  * revenue FY2025 = 1,234.5m everywhere EXCEPT the deck, which shows 1,334.5m  -> value_mismatch
  * net income FY2025 is blank in the XLSX                                     -> missing_item
  * capex FY2025 = 80.0m vs audited FY2024 baseline 50.0m (+60%)              -> baseline_anomaly
Everything else agrees (after rounding).
"""

import csv
from pathlib import Path

from openpyxl import Workbook
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfgen import canvas

COMPANY = {
    "id": "T001",
    "name": "Testco Holdings Inc.",
    "sector": "Industrial",
    "currency": "USD",
    "fiscal_year_end": "31 December",
    "segments": ["Machinery", "Services"],
}
FY = 2025

# True values in USD
TRUTH = {
    ("revenue", "FY2025"): 1_234_500_000,
    ("revenue", "FY2024"): 1_100_000_000,
    ("ebitda", "FY2025"): 210_300_000,
    ("ebitda", "FY2024"): 190_000_000,
    ("net_income", "FY2025"): 95_200_000,
    ("net_income", "FY2024"): 80_100_000,
    ("capex", "FY2025"): 80_000_000,
    ("capex", "FY2024"): 50_000_000,
    ("segment_revenue:Machinery", "FY2025"): 800_000_000,
    ("segment_revenue:Machinery", "FY2024"): 700_000_000,
    ("segment_revenue:Services", "FY2025"): 434_500_000,
    ("segment_revenue:Services", "FY2024"): 400_000_000,
}
BASELINE = [  # audited history (FY2024 and earlier)
    ("revenue", "FY2024", 1_100_000_000),
    ("ebitda", "FY2024", 190_000_000),
    ("net_income", "FY2024", 80_100_000),
    ("capex", "FY2024", 50_000_000),
    ("revenue", "FY2023", 1_000_000_000),
    ("capex", "FY2023", 48_000_000),
]


def _m(v: float) -> str:
    return f"{v / 1e6:,.1f}"


def _table(c, x, y, title, header, rows):
    c.setFont("Helvetica-Bold", 10)
    c.drawString(x, y, title)
    y -= 16
    c.setFont("Helvetica-Bold", 9)
    c.drawString(x, y, header[0])
    for j, h in enumerate(header[1:]):
        c.drawRightString(x + 330 + 90 * j, y, h)
    y -= 14
    c.setFont("Helvetica", 9)
    for label, vals in rows:
        c.drawString(x, y, label)
        for j, v in enumerate(vals):
            c.drawRightString(x + 330 + 90 * j, y, v)
        y -= 14
    return y - 20


def build(folder: Path) -> list[Path]:
    folder.mkdir(parents=True, exist_ok=True)
    t = TRUTH

    # Annual report: page 1 narrative, page 2 tables.
    ar = folder / "annual_report.pdf"
    c = canvas.Canvas(str(ar), pagesize=A4)
    c.drawString(50, 780, "Testco Holdings Inc. Annual Report FY2025")
    c.drawString(50, 700, "Narrative text without tables.")
    c.showPage()
    y = _table(
        c,
        50,
        760,
        "Consolidated statement of profit or loss (USD million)",
        ["", "FY2025", "FY2024"],
        [
            ("Revenue from operations", [_m(t[("revenue", "FY2025")]), _m(t[("revenue", "FY2024")])]),
            ("EBITDA", [_m(t[("ebitda", "FY2025")]), _m(t[("ebitda", "FY2024")])]),
            ("Profit for the year", [_m(t[("net_income", "FY2025")]), _m(t[("net_income", "FY2024")])]),
        ],
    )
    y = _table(
        c,
        50,
        y,
        "Note 4: Operating segments (USD million)",
        ["Segment", "FY2025", "FY2024"],
        [
            (
                "Machinery",
                [_m(t[("segment_revenue:Machinery", "FY2025")]), _m(t[("segment_revenue:Machinery", "FY2024")])],
            ),
            (
                "Services",
                [_m(t[("segment_revenue:Services", "FY2025")]), _m(t[("segment_revenue:Services", "FY2024")])],
            ),
            ("Total", [_m(t[("revenue", "FY2025")]), _m(t[("revenue", "FY2024")])]),
        ],
    )
    _table(
        c,
        50,
        y,
        "Note 12: Property, plant and equipment (USD million)",
        ["", "FY2025", "FY2024"],
        [
            ("Additions to property, plant and equipment", [_m(t[("capex", "FY2025")]), _m(t[("capex", "FY2024")])]),
        ],
    )
    c.showPage()
    c.save()

    # Board deck with the injected revenue error.
    deck = folder / "board_deck.pdf"
    c = canvas.Canvas(str(deck), pagesize=landscape(A4))
    _table(
        c,
        50,
        520,
        "Key metrics (USD million)",
        ["Metric", "FY2025", "FY2024", "YoY"],
        [
            ("Net sales", ["1,334.5", _m(t[("revenue", "FY2024")]), "+21.3%"]),
            ("EBITDA", [_m(t[("ebitda", "FY2025")]), _m(t[("ebitda", "FY2024")]), "+10.7%"]),
            ("Capex", [_m(t[("capex", "FY2025")]), _m(t[("capex", "FY2024")]), "+60.0%"]),
        ],
    )
    c.showPage()
    c.save()

    # Management accounts in USD '000, with net income FY2025 left blank.
    xlsx = folder / "management_accounts.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "P&L"
    ws["A1"] = "Testco - Income statement"
    ws["A2"] = "All amounts in USD '000"
    ws.append([])
    ws.append(["Line item", "FY25", "FY24"])
    for label, metric in [
        ("Total revenue", "revenue"),
        ("EBITDA", "ebitda"),
        ("Net income", "net_income"),
        ("Capital expenditure", "capex"),
    ]:
        cur = None if metric == "net_income" else round(t[(metric, "FY2025")] / 1000)
        ws.append([label, cur, round(t[(metric, "FY2024")] / 1000)])
    for row in ws.iter_rows(min_row=5, min_col=2, max_col=3):
        for cell in row:
            cell.number_format = "#,##0;(#,##0)"
    wb.save(xlsx)

    # Long-format KPI CSV in whole dollars.
    kpis = folder / "quarterly_kpis.csv"
    with open(kpis, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["entity_code", "period", "metric_code", "value", "currency", "unit"])
        codes = [
            ("REV", "revenue"),
            ("NI", "net_income"),
            ("CAPEX", "capex"),
            ("SEG_MACHINERY", "segment_revenue:Machinery"),
            ("SEG_SERVICES", "segment_revenue:Services"),
        ]
        for code, metric in codes:
            w.writerow(["T001", "FY2025", code, t[(metric, "FY2025")], "USD", "units"])
    return [ar, deck, xlsx, kpis]
